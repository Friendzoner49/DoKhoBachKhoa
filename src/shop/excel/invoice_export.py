"""Export one invoice to Excel for printing/sending to the customer.
Line totals use real formulas so the file can still be tweaked by hand
without the numbers going stale (the app itself stays the source of
truth for debt/profit — this file is for display only, see CLAUDE.md)."""
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from shop.core.customers import get_customer
from shop.core.invoices import get_invoice

HEADERS = ["Mã hàng", "Tên mặt hàng", "Đơn vị", "Số lượng", "Đơn giá (RUB)", "Thành tiền (RUB)"]


def export_invoice(conn: sqlite3.Connection, invoice_id: int, out_path: Path) -> Path:
    invoice = get_invoice(conn, invoice_id)
    if invoice is None:
        raise ValueError(f"Unknown invoice id: {invoice_id}")
    customer = get_customer(conn, invoice["customer_id"])

    wb = Workbook()
    ws = wb.active
    ws.title = f"HD {invoice_id}"

    ws.append([f"HÓA ĐƠN #{invoice_id}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Khách hàng: {customer['name']}"])
    if customer.get("address"):
        ws.append([f"Địa chỉ: {customer['address']}"])
    if customer.get("contact"):
        ws.append([f"Liên hệ: {customer['contact']}"])
    ws.append([f"Ngày: {invoice['invoice_date']}"])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(HEADERS)
    for cell in ws[header_row]:
        cell.font = Font(bold=True)

    first_data_row = header_row + 1
    for line in invoice["lines"]:
        ws.append([
            line["product_code"],
            line["product_name_snapshot"],
            line["unit_snapshot"],
            line["qty"],
            line["unit_price"],
            None,  # filled with a formula below
        ])
        row = ws.max_row
        # ROUND(...,0) matches shop.core.pricing.round_rub (half-up for
        # positive values), so this file's totals agree with the debt
        # the app recorded instead of showing fractional RUB.
        ws.cell(row=row, column=6, value=f"=ROUND(D{row}*E{row},0)")
    last_data_row = ws.max_row

    total_row = last_data_row + 1
    ws.cell(row=total_row, column=5, value="TỔNG CỘNG").font = Font(bold=True)
    ws.cell(row=total_row, column=6, value=f"=SUM(F{first_data_row}:F{last_data_row})").font = Font(bold=True)

    widths = [10, 40, 10, 10, 14, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = w

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
