"""Export the catalog to an Excel file shaped like the shop's own
original price list (ID, TEN MAT HANG, DON VI, DON GIA GOC, DON GIA,
SLL, GIA SLL, TRANG THAI), for the owner to browse/print — not the
public site export, which omits cost price entirely.
"""
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from shop.core.pricing import PUBLIC_MARKUP_PCT, bulk_price, unit_price
from shop.core.products import list_products

HEADERS = ["ID", "TÊN MẶT HÀNG", "ĐƠN VỊ", "ĐƠN GIÁ GỐC", "ĐƠN GIÁ", "SLL", "GIÁ SLL", "TRẠNG THÁI"]


def export_catalog(conn: sqlite3.Connection, out_path: Path, markup_pct=PUBLIC_MARKUP_PCT) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh muc"
    ws.append(HEADERS)

    for p in list_products(conn, active_only=False):
        price = unit_price(p["cost_price"], markup_pct)
        bulk = bulk_price(price, p["bulk_qty"]) if p["bulk_qty"] else 0
        ws.append([
            p["code"],
            p["name"],
            p["unit"],
            p["cost_price"],
            price,
            p["bulk_qty"] or 0,
            bulk,
            "Đang bán" if p["active"] else "Ngừng bán",
        ])

    for i, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(header) + 2)
    ws.column_dimensions["B"].width = 45

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
